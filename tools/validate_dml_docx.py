from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document


base = Path(r"C:\dev_projects\OY_dml\Kurumsal\Word_Dokumanlari")

for path in sorted(base.glob("*.docx")):
    try:
        with ZipFile(path) as zf:
            names = set(zf.namelist())
            assert "[Content_Types].xml" in names
            assert "word/document.xml" in names
        doc = Document(path)
        nonempty = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        print(
            f"{path.name}: OK | paragraphs={len(nonempty)} | "
            f"tables={len(doc.tables)} | size={path.stat().st_size}"
        )
        print(f"  title={nonempty[0][:100] if nonempty else 'EMPTY'}")
    except (BadZipFile, AssertionError, Exception) as exc:
        print(f"{path.name}: ERROR {exc}")
