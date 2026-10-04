from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(r"C:\dev_projects\OY_dml\Kurumsal")
OUT = BASE / "Word_Dokumanlari"


FILES = [
    "01_DML_Arastirma_Merkezi_Yonergesi_Taslagi.md",
    "02_DML_Mevcut_Durum_ve_Envanter_Raporu_Taslagi.md",
    "03_DML_Kullanim_Alanlari_ve_Hizmet_Kapsami_Haritasi_Taslagi.md",
    "04_DML_Organizasyon_Gorev_Dagilimi_ve_Isleyis_Plani_Taslagi.md",
    "05_DML_Kullanim_Guvenlik_ve_Rezervasyon_Prosedurleri_Taslagi.md",
    "06_DML_Kurumsal_Gelisim_Tanitim_ve_Egitim_Plani_Taslagi.md",
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = margins.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:cs"), "Aptos")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_formatted_text(paragraph, text: str, size: float | None = None) -> None:
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(part)
            set_font(run, size=size)


def apply_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

    for name, size, before, after in [
        ("Title", 18, 0, 12),
        ("Heading 1", 15, 14, 6),
        ("Heading 2", 12.5, 10, 4),
        ("Heading 3", 11, 8, 3),
    ]:
        style = doc.styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx].cells
        for c_idx, value in enumerate(row):
            cell = cells[c_idx]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, "1F4E79")
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "F4F8FC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_formatted_text(p, value.strip(), size=9)
            for run in p.runs:
                if r_idx == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
                elif len(value.strip()) <= 18 and c_idx != len(row) - 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    table_lines = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        raw = lines[idx].strip()
        if not re.fullmatch(r"\|[\s:\-|\u2013]+\|?", raw):
            table_lines.append(raw)
        idx += 1
    rows = []
    for raw in table_lines:
        cells = [cell.strip() for cell in raw.strip("|").split("|")]
        rows.append(cells)
    return rows, idx


def add_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    add_formatted_text(p, text)


def build_docx(md_path: Path) -> Path:
    text = md_path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    apply_styles(doc)

    idx = 0
    while idx < len(lines):
        line = lines[idx].strip()
        if not line:
            idx += 1
            continue
        if line.startswith("|"):
            rows, idx = parse_table(lines, idx)
            add_table(doc, rows)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_formatted_text(p, line[2:].strip(), size=18)
        elif line.startswith("## "):
            add_paragraph(doc, line[3:].strip(), "Heading 1")
        elif line.startswith("### "):
            add_paragraph(doc, line[4:].strip(), "Heading 2")
        elif re.match(r"^\d+\.\s+", line):
            add_paragraph(doc, line, "List Number")
        elif line.startswith("- "):
            add_paragraph(doc, line[2:].strip(), "List Bullet")
        else:
            add_paragraph(doc, line)
        idx += 1

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Taslak belge - üniversite görüşmeleriyle geliştirilecektir")
        set_font(run, size=8)

    OUT.mkdir(parents=True, exist_ok=True)
    out_path = OUT / (md_path.stem + ".docx")
    doc.save(out_path)
    return out_path


def main() -> None:
    outputs = [build_docx(BASE / name) for name in FILES]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
